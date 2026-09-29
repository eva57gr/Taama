"""Turns an input (raw text, PDF, image, Word, Excel, HTML file or URL) into laid-out text blocks with provenance."""

from __future__ import annotations

import logging
import re
import statistics
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

from models import Block, Provenance
import ocr
from ocr import ocr_lines
from rulebank import REPO_ROOT, squash

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
OFFICE_SUFFIXES = {".docx", ".xlsx"}
log = logging.getLogger("claimcheck.ingest")

# Words OCR must never split and that help re-split run-together text (ingredient / technology names).
DOMAIN_WORDS = {
    "adaptogen", "adaptogens", "adaptogenic", "nootropic", "nootropics", "methylliberine", "theacrine", "teacrine",
    "sowell", "cereboost", "viacap", "postbiotic", "postbiotics", "prebiotic", "prebiotics", "enzogenol",
    "neuroberry", "suntheanine", "theanine", "erythritol", "flavonoids", "flavonoid", "polyphenols", "polyphenol",
    "proanthocyanidins", "methylglyoxal", "manuka", "glycosides", "microbiome", "ginsenosides", "isothiocyanates",
    "menaquinone", "ubiquinone", "phytosome", "pyrroloquinoline", "neuroscientists", "neuroscientist", "kucha",
    "oligomeric", "bioenergetic", "hypromellose", "adenosylcobalamin", "methylcobalamin", "magnafolate",
}
# Function words OCR glues to a neighbour and that are too short for the segmenter.
GLUED = {"anda": "and a", "witha": "with a", "fora": "for a", "ina": "in a", "isa": "is a", "toa": "to a", "upto": "up to"}
UNITS = {"mg", "g", "kg", "mcg", "ml", "l", "hr", "hrs", "h", "min", "mins", "kj", "cal", "kcal", "x", "s",
         "am", "pm", "st", "nd", "rd", "th", "d", "mm", "cm", "oz", "dfe", "iu"}
CONNECTORS = {"and", "or", "for", "to", "of", "with", "a", "an", "the", "that", "in", "by", "from", "your",
              "&", "+", "-", "is", "are", "as", "on", "at", "our", "into", "than", "which", "who", "plus"}


def rel(path: Path) -> str:
    try:
        return Path(path).resolve().relative_to(REPO_ROOT).as_posix()
    except ValueError:
        return str(path)


# ---------------------------------------------------------------------------------------------- text


def from_text(text: str, name: str = "<text>") -> list[Block]:
    """Raw text: blank lines separate blocks; single line breaks inside a block are kept as bullets."""
    blocks = []
    for i, para in enumerate(re.split(r"\n\s*\n", text.strip())):
        lines = [l.strip(" \t•·*-–") for l in para.splitlines() if l.strip()]
        if not lines:
            continue
        # A paragraph whose lines wrap (no terminal punctuation, next starts lowercase) is one block.
        merged = [lines[0]]
        for line in lines[1:]:
            if not re.search(r"[.!?:]$", merged[-1]) and (line[:1].islower() or merged[-1].split()[-1].lower() in CONNECTORS):
                merged[-1] += " " + line
            else:
                merged.append(line)
        for j, line in enumerate(merged):
            blocks.append(Block(squash(line), Provenance(name, f"paragraph {i + 1}, line {j + 1}", "text")))
    return blocks


# ----------------------------------------------------------------------------------------------- PDF


def from_pdf(path: Path) -> list[Block]:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams, LTTextBox

    blocks = []
    for page_no, page in enumerate(extract_pages(str(path), laparams=LAParams()), start=1):
        boxes = [el for el in page if isinstance(el, LTTextBox)]
        boxes.sort(key=lambda b: (-round(b.y1), round(b.x0)))
        for n, box in enumerate(boxes, start=1):
            lines = [l.get_text().strip() for l in box if l.get_text().strip()]
            if not lines:
                continue
            # Rotated / stacked single characters (barcodes, vertical legal lines) are not claim text.
            if sum(len(l) <= 2 for l in lines) > len(lines) / 2:
                continue
            text = squash(" ".join(lines))
            loc = f"page {page_no}, text box {n} at ({box.x0:.0f}, {page.height - box.y1:.0f})"
            blocks.append(Block(text, Provenance(rel(path), loc, "pdf-text-layer")))
    return blocks


# --------------------------------------------------------------------------------------------- images


def _load_wordninja():
    try:
        import wordninja
    except ImportError:  # pragma: no cover
        return None, set()
    return wordninja, set(wordninja.DEFAULT_LANGUAGE_MODEL._wordcost)


def _split_run(run: str, wordninja, vocab: set[str], protected: set[str]) -> str:
    low = run.lower()
    if low in vocab or low in protected or wordninja is None:
        return run
    for word in sorted(protected, key=len, reverse=True):
        i = low.find(word)
        if i >= 0 and len(word) >= 5:
            left, mid, right = run[:i], run[i:i + len(word)], run[i + len(word):]
            parts = [_split_run(left, wordninja, vocab, protected) if left else "", mid,
                     _split_run(right, wordninja, vocab, protected) if right else ""]
            return " ".join(p for p in parts if p)
    return " ".join(wordninja.split(run))


def respace(text: str, protected: set[str] = frozenset()) -> str:
    """Re-inserts spaces OCR drops on tightly kerned artwork ("Protectskeystructuresofthebrain")."""
    wordninja, vocab = _load_wordninja()
    protected = DOMAIN_WORDS | {w.lower() for w in protected}
    t = text
    t = re.sub(r"\b([A-Z][A-Za-z]{2,})\?(?=[\s,.)\]]|$)", r"\1®", t)         # Neuroberry? -> Neuroberry® (OCR)
    t = re.sub(r"(?<=[A-Z]-)O(?=\d)", "0", t)                                # AM-O2 -> AM-02
    t = re.sub(r"(?<=[a-z]{2})(?=[$\d])", " ", t)                           # in1hr -> in 1hr, than$1 -> than $1
    t = re.sub(r"(?<=\d)([A-Za-z]+)", lambda m: m.group(1) if m.group(1).lower() in UNITS else " " + m.group(1), t)
    t = re.sub(r"(?<=%)(?=[A-Za-z])", " ", t)                               # 100%natural
    t = re.sub(r"(?<=[A-Za-z])&(?=[A-Za-z])", " & ", t)                     # Designed&tested
    t = re.sub(r"(?<=[a-z]),(?=[A-Za-z])", ", ", t)                          # clearer,faster
    t = re.sub(r"(?<=[a-z0-9])\.(?=[A-Z])", ". ", t)                         # daily.Safe
    t = re.sub(r"(?<=[a-z]{2}):(?=[A-Za-z])", ": ", t)                       # Food:Patented
    t = re.sub(r"(?<=[a-z]{2})(?=[A-Z]{2})", " ", t)                          # ofAM-O2
    def _article(m):                                                          # Afast-acting
        whole, rest = ("a" + m.group(1)).split("-")[0], m.group(1).split("-")[0]
        glued = whole not in vocab and whole not in protected and rest in vocab
        return "A " + m.group(1) if glued else "A" + m.group(1)
    t = re.sub(r"\bA([a-z][a-z-]{2,})", _article, t)
    t = re.sub(r"\b[a-z]{3,5}\b", lambda m: GLUED.get(m.group(0), m.group(0)), t)
    t = re.sub(r"[A-Za-z]{7,}", lambda m: _split_run(m.group(0), wordninja, vocab, protected), t)
    return squash(t)


def _overlap(a: list[float], b: list[float]) -> float:
    inter = min(a[2], b[2]) - max(a[0], b[0])
    return max(0.0, inter) / max(1.0, min(a[2] - a[0], b[2] - b[0]))


def group_lines(lines: list[dict]) -> list[list[dict]]:
    """Chains OCR lines into blocks: a line continues the one above it if it sits directly below,
    overlaps it horizontally, and reads as a continuation (wrap, lowercase start, connector, or a
    multi-line heading)."""
    if not lines:
        return []
    lines = sorted(lines, key=lambda l: (l["box"][1], l["box"][0]))
    heights = [l["box"][3] - l["box"][1] for l in lines]
    median_h = statistics.median(heights)
    nxt: dict[int, int] = {}
    taken: set[int] = set()
    for i, a in enumerate(lines):
        ah = a["box"][3] - a["box"][1]
        best = None
        for j in range(i + 1, len(lines)):
            b = lines[j]
            if j in taken or b["box"][1] < a["box"][1] + 0.5 * ah:
                continue
            bh = b["box"][3] - b["box"][1]
            gap = b["box"][1] - a["box"][3]
            if gap > 0.9 * max(ah, bh):
                continue
            if _overlap(a["box"], b["box"]) < 0.3 or not (0.55 <= bh / ah <= 1.8):
                continue
            best = j
            break
        if best is None:
            continue
        b = lines[best]
        at, bt = a["text"].strip(), b["text"].strip()
        last = at.split()[-1].lower() if at.split() else ""
        width_a = a["box"][2] - a["box"][0]
        column_w = max(l["box"][2] - l["box"][0] for l in lines if _overlap(l["box"], a["box"]) >= 0.3)
        ends_sentence = bool(re.search(r"[.!?:]$", at))
        heading = ah >= 1.4 * median_h and (b["box"][3] - b["box"][1]) >= 1.4 * median_h
        left_aligned = abs(a["box"][0] - b["box"][0]) <= 1.5 * median_h
        continues = (
            bt[:1].islower()
            or last in CONNECTORS
            or at.endswith((",", "-", "&", "+"))
            or (heading and not ends_sentence)
            or (left_aligned and width_a >= 0.8 * column_w and not ends_sentence and len(at) > 25)
        )
        if continues:
            nxt[i] = best
            taken.add(best)
    blocks, seen = [], set()
    for i in range(len(lines)):
        if i in seen or i in taken:
            continue
        chain, k = [], i
        while k is not None and k not in seen:
            chain.append(lines[k])
            seen.add(k)
            k = nxt.get(k)
        blocks.append(chain)
    return blocks


def from_image(path: Path, protected: set[str] = frozenset(), refresh_ocr: bool = False) -> list[Block]:
    engine = ocr.backend()
    log.info("image %s -> OCR engine %s%s", Path(path).name, engine, " (refresh)" if refresh_ocr else "")
    if engine == "openai":
        texts, method = ocr.openai_blocks(path, refresh=refresh_ocr)
        return [Block(squash(t), Provenance(rel(path), f"text block {i + 1}", method))
                for i, t in enumerate(texts) if squash(t)]
    lines, method = ocr_lines(path, refresh=refresh_ocr)
    log.info("RapidOCR %s: %d lines (%s)", Path(path).name, len(lines), method)
    lines = [dict(l, text=respace(l["text"], protected)) for l in lines]
    blocks = []
    for chain in group_lines(lines):
        text = squash(" ".join(l["text"] for l in chain))
        x0 = min(l["box"][0] for l in chain)
        y0 = min(l["box"][1] for l in chain)
        x1 = max(l["box"][2] for l in chain)
        y1 = max(l["box"][3] for l in chain)
        conf = min(l["score"] for l in chain)
        loc = f"region ({x0:.0f}, {y0:.0f})-({x1:.0f}, {y1:.0f}) px, {len(chain)} line(s), min OCR confidence {conf:.2f}"
        blocks.append(Block(text, Provenance(rel(path), loc, method), (x0, y0, x1, y1), conf))
    return blocks


# -------------------------------------------------------------------------------------- Word / Excel


def from_docx(path: Path) -> list[Block]:
    from docx import Document

    doc = Document(str(path))
    blocks = []
    for i, p in enumerate(doc.paragraphs, start=1):
        if squash(p.text):
            blocks.append(Block(squash(p.text), Provenance(rel(path), f"paragraph {i}", "docx")))
    for t, table in enumerate(doc.tables, start=1):
        for r, row in enumerate(table.rows, start=1):
            seen = []
            for cell in row.cells:  # merged cells repeat; keep each once
                text = squash(cell.text)
                if text and text not in seen:
                    seen.append(text)
            for c, text in enumerate(seen, start=1):
                blocks.append(Block(text, Provenance(rel(path), f"table {t}, row {r}, cell {c}", "docx")))
    return blocks


def from_xlsx(path: Path) -> list[Block]:
    import zipfile
    import zlib

    from openpyxl import load_workbook

    blocks = []
    try:
        with open(path, "rb") as fh:
            wb = load_workbook(fh, read_only=True, data_only=True)
            try:
                for ws in wb.worksheets:
                    for row in ws.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str) and squash(cell.value):
                                blocks.append(Block(squash(cell.value), Provenance(
                                    rel(path), f"sheet {ws.title!r}, cell {cell.coordinate}", "xlsx")))
            finally:
                wb.close()
    except (zipfile.BadZipFile, zlib.error, KeyError, OSError):
        blocks = _salvage_xlsx(path)
        if not blocks:
            raise
    return blocks


def _salvage_xlsx(path: Path) -> list[Block]:
    """Read text cells straight from the worksheet XML when the workbook itself is damaged."""
    import re
    import struct
    import zlib
    import xml.etree.ElementTree as ET

    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

    def text_of(el) -> str:
        return "".join(t.text or "" for t in el.iter(f"{ns}t"))

    # The central directory may be broken, so walk the local file headers instead.
    data = path.read_bytes()
    parts: dict[str, bytes] = {}
    for m in re.finditer(rb"PK\x03\x04", data):
        o = m.start()
        if o + 30 > len(data):
            break
        flags, method, _, _, _, csize, _, nlen, xlen = struct.unpack("<HHHHIIIHH", data[o + 6:o + 30])
        name = data[o + 30:o + 30 + nlen].decode("utf-8", "replace")
        start = o + 30 + nlen + xlen
        if flags & 0x08 or not csize:  # sizes stored after the data; skip
            continue
        raw = data[start:start + csize]
        try:
            parts[name] = zlib.decompress(raw, -15) if method == 8 else raw if method == 0 else b""
        except zlib.error:
            continue

    def read(name):
        try:
            return ET.fromstring(parts[name])
        except (KeyError, ET.ParseError):
            return None

    blocks = []
    sst = read("xl/sharedStrings.xml")
    shared = [text_of(si) for si in sst.iter(f"{ns}si")] if sst is not None else []
    sheets = sorted((n for n in parts if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)),
                    key=lambda n: int(re.search(r"\d+", n.rsplit("/", 1)[1]).group()))
    for name in sheets:
        root = read(name)
        if root is None:
            continue
        sheet = name.rsplit("/", 1)[1].removesuffix(".xml")
        for c in root.iter(f"{ns}c"):
            kind, v = c.get("t"), c.find(f"{ns}v")
            if kind == "inlineStr":
                value = text_of(c)
            elif kind == "s" and v is not None and (v.text or "").isdigit() and int(v.text) < len(shared):
                value = shared[int(v.text)]
            elif kind == "str" and v is not None:
                value = v.text or ""
            else:
                continue
            if squash(value):
                blocks.append(Block(squash(value), Provenance(
                    rel(path), f"{sheet} (recovered), cell {c.get('r')}", "xlsx")))
    return blocks


# ----------------------------------------------------------------------------------------- HTML / URL


class _HTMLBlocks(HTMLParser):
    BLOCK = {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "td", "th", "dt", "dd", "div", "section", "br", "tr", "blockquote", "figcaption", "span"}
    SKIP = {"script", "style", "noscript", "svg", "head", "nav", "footer"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[str] = []
        self._buf: list[str] = []
        self._skip = 0

    def _flush(self):
        text = squash("".join(self._buf))
        if text:
            self.blocks.append(text)
        self._buf = []

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK - {"span"}:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag in self.BLOCK - {"span"}:
            self._flush()

    def handle_data(self, data):
        if not self._skip:
            self._buf.append(data)

    def close(self):
        super().close()
        self._flush()


def from_html(html: str, name: str) -> list[Block]:
    parser = _HTMLBlocks()
    parser.feed(html)
    parser.close()
    return [Block(t, Provenance(name, f"html block {i + 1}", "html")) for i, t in enumerate(parser.blocks)]


def from_url(url: str) -> list[Block]:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (claimcheck)"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        html = resp.read().decode(resp.headers.get_content_charset() or "utf-8", errors="replace")
    return from_html(html, url)


# ---------------------------------------------------------------------------------------------- entry


def ingest(source: str | Path, protected: set[str] = frozenset(), refresh_ocr: bool = False) -> list[Block]:
    s = str(source)
    if re.match(r"https?://", s):
        return from_url(s)
    path = Path(s)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return from_pdf(path)
    if suffix in IMAGE_SUFFIXES:
        return from_image(path, protected, refresh_ocr)
    if suffix in OFFICE_SUFFIXES:
        import zipfile
        import zlib

        try:
            return from_docx(path) if suffix == ".docx" else from_xlsx(path)
        except (zipfile.BadZipFile, zlib.error, OSError, KeyError) as e:
            raise ValueError(f"cannot read {path.name}: the {suffix} file is damaged or not a valid Office file") from e
    if suffix in {".html", ".htm"}:
        return from_html(path.read_text(encoding="utf-8", errors="replace"), rel(path))
    if suffix in {".txt", ".md"}:
        return from_text(path.read_text(encoding="utf-8"), rel(path))
    raise ValueError(f"unsupported input type: {s}")
